import re
from pathlib import Path

import openpyxl
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor as DocxRGBColor
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt as PPt

DARK = RGBColor(0x1E, 0x2A, 0x5A)
ACCENT = RGBColor(0xF2, 0xC9, 0x4C)
GRAY = RGBColor(0x6B, 0x6B, 0x6B)

COVER_LABEL = {"uz": "REFERAT", "ru": "РЕФЕРАТ", "en": "REPORT"}
PREPARED = {
    "uz": {"by": "Bajaruvchi:", "check": "Tekshiruvchi:"},
    "ru": {"by": "Выполнил:", "check": "Проверил:"},
    "en": {"by": "Prepared by:", "check": "Checked by:"},
}
THANKS = {"uz": "Eʼtiboringiz uchun rahmat!", "ru": "Спасибо за внимание!", "en": "Thank you for your attention!"}


def _set_font(run, size=14, bold=False, name="Times New Roman"):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), name)


def _page_number(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)
    _set_font(run, size=12)


def build_referat(data: dict, lang: str, path: Path) -> Path:
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin, sec.bottom_margin = Cm(2), Cm(2)
    sec.left_margin, sec.right_margin = Cm(3), Cm(1.5)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(14)
    pf = normal.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE

    footer_p = sec.footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _page_number(footer_p)

    for _ in range(4):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_font(p.add_run(COVER_LABEL.get(lang, "REFERAT")), size=16, bold=True)
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_font(p.add_run(data.get("title", "")), size=14, bold=True)
    for _ in range(6):
        doc.add_paragraph()
    labels = PREPARED.get(lang, PREPARED["uz"])
    for key in ("by", "check"):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        _set_font(p.add_run(f"{labels[key]} {'.' * 30}"), size=14)

    doc.add_page_break()

    for i, section in enumerate(data.get("sections", [])):
        heading = str(section.get("heading", "")).strip()
        if heading:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(6)
            _set_font(p.add_run(heading), size=14, bold=True)
        for para in section.get("paragraphs", []):
            text = str(para).strip()
            if not text:
                continue
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.first_line_indent = Cm(1.25)
            p.paragraph_format.space_after = Pt(6)
            _set_font(p.add_run(text), size=14)
        if i == len(data.get("sections", [])) - 1:
            doc.add_paragraph()

    doc.save(path)
    return path


def _add_box(slide, x, y, w, h, color):
    from pptx.enum.shapes import MSO_SHAPE
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def build_pptx(data: dict, lang: str, path: Path) -> Path:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    W, H = prs.slide_width, prs.slide_height

    slides = list(data.get("slides", []))
    if not slides:
        slides = [{"title": data.get("title", ""), "bullets": [], "notes": ""}]

    for idx, s in enumerate(slides):
        slide = prs.slides.add_slide(blank)
        title = re.sub(r"^\s*\d+[.)]\s*", "", str(s.get("title", "")).strip())
        bullets = [re.sub(r"^\s*\d+[.)]\s*", "", str(b).strip())
                   for b in s.get("bullets", []) if str(b).strip()]
        notes = str(s.get("notes", "")).strip()
        total_len = sum(len(b) for b in bullets)
        body_size = 20 if total_len <= 380 else (17 if total_len <= 550 else 15)

        if idx == 0:
            _add_box(slide, 0, 0, W, H, DARK)
            _add_box(slide, 0, H - Inches(0.35), W, Inches(0.35), ACCENT)
            tb = slide.shapes.add_textbox(Inches(1), Inches(2.4), W - Inches(2), Inches(2.2))
            tf = tb.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = 1
            r = p.add_run()
            r.text = str(data.get("title", title))
            r.font.size = PPt(44)
            r.font.bold = True
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            r.font.name = "Arial"
            sub = str(data.get("subtitle", "")).strip()
            if sub:
                p2 = tf.add_paragraph()
                p2.alignment = 1
                p2.space_before = PPt(18)
                r2 = p2.add_run()
                r2.text = sub
                r2.font.size = PPt(20)
                r2.font.color.rgb = ACCENT
                r2.font.name = "Arial"
        else:
            _add_box(slide, 0, 0, W, Inches(1.15), DARK)
            _add_box(slide, 0, Inches(1.15), W, Inches(0.08), ACCENT)
            tb = slide.shapes.add_textbox(Inches(0.6), Inches(0.18), W - Inches(1.2), Inches(0.85))
            tf = tb.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            r = p.add_run()
            r.text = title
            r.font.size = PPt(30)
            r.font.bold = True
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            r.font.name = "Arial"

            body = slide.shapes.add_textbox(Inches(0.7), Inches(1.7), W - Inches(1.6), H - Inches(2.6))
            tf = body.text_frame
            tf.word_wrap = True
            for bi, b in enumerate(bullets):
                p = tf.paragraphs[0] if bi == 0 else tf.add_paragraph()
                p.space_after = PPt(8) if body_size < 20 else PPt(12)
                glyph = p.add_run()
                glyph.text = "▸  "
                glyph.font.size = PPt(body_size)
                glyph.font.bold = True
                glyph.font.color.rgb = ACCENT
                glyph.font.name = "Arial"
                r = p.add_run()
                r.text = b
                r.font.size = PPt(body_size)
                r.font.color.rgb = RGBColor(0x2B, 0x2B, 0x2B)
                r.font.name = "Arial"
            if not bullets and notes:
                p = tf.paragraphs[0]
                r = p.add_run()
                r.text = notes
                r.font.size = PPt(18)
                r.font.color.rgb = GRAY
                r.font.name = "Arial"

            num = slide.shapes.add_textbox(W - Inches(1.2), H - Inches(0.55), Inches(0.9), Inches(0.4))
            p = num.text_frame.paragraphs[0]
            p.alignment = 2
            r = p.add_run()
            r.text = str(idx + 1)
            r.font.size = PPt(14)
            r.font.color.rgb = GRAY
            r.font.name = "Arial"

            ft = slide.shapes.add_textbox(Inches(0.4), H - Inches(0.55), Inches(4), Inches(0.4))
            p = ft.text_frame.paragraphs[0]
            r = p.add_run()
            r.text = THANKS.get(lang, "") if idx == len(slides) - 1 else ""
            r.font.size = PPt(14)
            r.font.color.rgb = GRAY
            r.font.name = "Arial"

        if notes:
            slide.notes_slide.notes_text_frame.text = notes

    prs.save(path)
    return path


def build_xlsx(data: dict, lang: str, path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Jadval"

    headers = [str(h) for h in data.get("headers", [])] or ["№", "Maʼlumot"]
    rows = [[("" if c is None else c) for c in row] for row in data.get("rows", [])]
    if not rows:
        rows = [["" for _ in headers]]

    dark_fill = PatternFill("solid", fgColor="1E2A5A")
    alt_fill = PatternFill("solid", fgColor="EEF1F8")
    white_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    body_font = Font(name="Arial", size=11)
    thin = Side(style="thin", color="B0B0B0")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)

    ncols = len(headers)
    last_col = ws.cell(row=1, column=ncols).column_letter
    ws.merge_cells(f"A1:{last_col}1")
    title_cell = ws["A1"]
    title_cell.value = data.get("title", "")
    title_cell.font = Font(name="Arial", size=13, bold=True, color="1E2A5A")
    title_cell.alignment = center
    ws.row_dimensions[1].height = 26

    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=2, column=c, value=h)
        cell.fill = dark_fill
        cell.font = white_font
        cell.alignment = center
        cell.border = border
    ws.row_dimensions[2].height = 22

    for r, row in enumerate(rows, 3):
        for c in range(1, ncols + 1):
            val = row[c - 1] if c - 1 < len(row) else ""
            cell = ws.cell(row=r, column=c, value=val)
            cell.font = body_font
            cell.alignment = left if c == 1 or not isinstance(val, (int, float)) else center
            cell.border = border
            if r % 2 == 1:
                cell.fill = alt_fill

    for c, h in enumerate(headers, 1):
        longest = max([len(str(h))] + [len(str(row[c - 1])) if c - 1 < len(row) else 0 for row in rows])
        ws.column_dimensions[ws.cell(row=2, column=c).column_letter].width = min(max(longest + 4, 10), 55)

    ws.freeze_panes = "A3"
    note = str(data.get("note", "") or "").strip()
    if note:
        ws.cell(row=len(rows) + 4, column=1, value=note).font = Font(name="Arial", size=9, italic=True, color="777777")

    num_col = next((c for c in range(1, ncols + 1)
                    if rows and all(isinstance(r[c - 1], (int, float)) for r in rows)), None)
    category_col = next((c for c in range(1, ncols + 1)
                         if c != num_col and rows and all(str(r[c - 1]).strip() for r in rows)), 1)
    if num_col and rows:
        from openpyxl.chart import BarChart, Reference
        chart = BarChart()
        chart.type = "col"
        chart.style = 10
        chart.title = data.get("title", "Grafik")
        chart.y_axis.title = headers[num_col - 1]
        data_ref = Reference(ws, min_col=num_col, min_row=2, max_row=len(rows) + 2)
        cats = Reference(ws, min_col=category_col, min_row=3, max_row=len(rows) + 2)
        chart.add_data(data_ref, titles_from_data=True)
        chart.set_categories(cats)
        chart.legend = None
        ws.add_chart(chart, f"{get_column_letter(ncols + 2)}2")

    wb.save(path)
    return path


def build_test_docx(questions: list[dict], lang: str, path: Path) -> Path:
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = sec.bottom_margin = Cm(2)
    sec.left_margin = sec.right_margin = Cm(2.5)
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(13)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_font(p.add_run("TEST SAVOLLARI"), size=16, bold=True)
    doc.add_paragraph()

    for i, q in enumerate(questions, 1):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        _set_font(p.add_run(f"{i}. {q.get('q', '')}"), size=13, bold=True)
        letters = "ABCDEFGH"
        for oi, opt in enumerate(q.get("options", [])):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.6)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            _set_font(p.add_run(f"{letters[oi]}) {opt}"), size=13)
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.6)
        p.paragraph_format.space_after = Pt(4)
        ans_idx = q.get("answer")
        ans_letter = letters[ans_idx] if isinstance(ans_idx, int) and 0 <= ans_idx < len(q.get("options", [])) else "?"
        color = DocxRGBColor(0x19, 0x7B, 0x30) if ans_letter != "?" else DocxRGBColor(0xC0, 0x39, 0x2B)
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.6)
        p.paragraph_format.space_after = Pt(4)
        _set_font(p.add_run(f"Toʻgʻri javob: {ans_letter}"), size=12)
        p.runs[0].font.color.rgb = color

    doc.save(path)
    return path


def build_text_docx(paragraphs: list[str], title: str, path: Path) -> Path:
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = sec.bottom_margin = Cm(2)
    sec.left_margin = sec.right_margin = Cm(2.5)
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    if title:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _set_font(p.add_run(title), size=14, bold=True)
        doc.add_paragraph()
    for para in paragraphs:
        if para.startswith("## "):
            p = doc.add_paragraph()
            _set_font(p.add_run(para[3:]), size=13, bold=True)
        else:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.first_line_indent = Cm(1)
            p.paragraph_format.space_after = Pt(4)
            _set_font(p.add_run(para), size=12)
    doc.save(path)
    return path
