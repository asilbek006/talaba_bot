import re
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
BADWORDS_PATHS = [
    ROOT / "data" / "badwords.xlsx",
    ROOT / "data" / "bad_words.txt",
    Path("/home/asilbek/Downloads/Telegram Desktop/badwords.xlsx"),
    Path("/home/asilbek/Downloads/badwords.xlsx"),
]


def _load_words() -> set[str]:
    words = set()
    for p in BADWORDS_PATHS:
        if not p.exists():
            continue
        try:
            if p.suffix == ".xlsx":
                from openpyxl import load_workbook

                wb = load_workbook(p, read_only=True, data_only=True)
                ws = wb.active
                for row in ws.iter_rows(values_only=True):
                    for v in row:
                        if not v:
                            continue
                        s = str(v).strip().lower()
                        if s and len(s) >= 2:
                            words.add(s)
            elif p.suffix == ".txt":
                for line in p.read_text(encoding="utf-8").splitlines():
                    s = line.strip().lower()
                    if s and not s.startswith("#") and len(s) >= 2:
                        words.add(s)
        except Exception:
            continue
    return words


_BADWORDS = _load_words()
_BADWORDS_SORTED = sorted(_BADWORDS, key=len, reverse=True)


def _matches(w: str, t: str) -> bool:
    return re.search(rf"\b{re.escape(w)}\b", t) is not None


def has_badword(text: str) -> bool:
    if not text:
        return False
    from services.moderation import is_insult

    if is_insult(text):
        return True
    t = text.lower()
    for w in _BADWORDS_SORTED:
        if _matches(w, t):
            return True
    return False


def get_badwords(text: str) -> list[str]:
    found = []
    if not text:
        return found
    t = text.lower()
    for w in _BADWORDS_SORTED:
        if _matches(w, t):
            found.append(w)
    seen = set()
    res = []
    for f in found:
        if f not in seen:
            res.append(f)
            seen.add(f)
    return res
