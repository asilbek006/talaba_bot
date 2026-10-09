import csv
import re
from pathlib import Path

from docx import Document

import config
from services.validation import validate_docx


def extract_docx(path):
    validate_docx(Path(path))
    doc = Document(path)
    lines = []
    for p in doc.paragraphs:
        text = p.text.strip()
        if text:
            lines.append(text)
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            lines.append(" | ".join(c for c in cells if c))
    if sum(map(len, lines)) > config.MAX_TEXT_CHARS:
        raise ValueError("Word matni juda katta")
    return lines


def extract_txt(path) -> list[str]:
    p = Path(path)
    raw = p.read_bytes()
    for enc in ("utf-8-sig", "utf-8", "cp1251", "latin-1"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        text = raw.decode("utf-8", errors="ignore")
    return [line.strip() for line in text.splitlines() if line.strip()]


ANSWER_MAP = {
    "A": 0, "B": 1, "C": 2, "D": 3, "E": 4, "F": 5, "G": 6, "H": 7,
    "А": 0, "Б": 1, "В": 2, "Г": 3, "Д": 4, "Е": 5,
    "1": 0, "2": 1, "3": 2, "4": 3, "5": 4, "6": 5,
}

Q_RE = re.compile(r"^\s*(savol|question|вопрос)\s*[:.\-)]*\s*", re.I)
NUM_RE = re.compile(r"^\s*\d+\s*[.)]\s*")
OPT_RE = re.compile(r"^\s*([\+\*\-]?\s*\[?([A-Ha-hА-Еа-е])\]?[\s.)\-:])\s*")
ANS_RE = re.compile(
    r"(?:to[ʻ'’`]?g[ʻ'’`]?ri\s*javob|javob[a-z]*|answer|ответ|прав[a-z]*\s*ответ)\s*[:.\-)]*\s*\**\s*([a-hA-Hа-еА-Е0-9])(?=\s*[)\s.]|$)",
    re.I,
)


def parse_questions(lines) -> list[dict]:
    qs = []
    cur_q = None
    for line in lines:
        s = line.strip()
        if not s:
            continue
        ans_match = ANS_RE.search(s)
        if ans_match:
            if cur_q:
                cur_q["answer"] = ANSWER_MAP.get(ans_match.group(1).upper())
            continue
        opt = OPT_RE.match(s)
        if cur_q and opt:
            # Check if option starts with + or * (marking correct answer)
            is_marked_correct = s.startswith("+") or s.startswith("*")
            cleaned_opt = OPT_RE.sub("", s, count=1).strip()
            if is_marked_correct and cur_q.get("answer") is None:
                cur_q["answer"] = len(cur_q["options"])
            cur_q["options"].append(cleaned_opt)
            continue
        if Q_RE.match(s) or NUM_RE.match(s) or not (cur_q and cur_q["options"]):
            if cur_q and cur_q["q"] and cur_q["options"]:
                qs.append(cur_q)
            cur_q = {"q": Q_RE.sub("", s, count=1).strip(), "options": [], "answer": None}
            cur_q["q"] = NUM_RE.sub("", cur_q["q"], count=1).strip()
        elif cur_q:
            cur_q["q"] += " " + s
    if cur_q and cur_q["q"] and cur_q["options"]:
        qs.append(cur_q)
    return [
        q
        for q in qs
        if 2 <= len(q["options"]) <= 8
        and type(q["answer"]) is int
        and 0 <= q["answer"] < len(q["options"])
        and len(q["q"]) <= 2500
    ][:500]


def parse_table_questions(rows) -> list[dict]:
    letters_map = {
        "a": 0, "b": 1, "c": 2, "d": 3, "e": 4, "f": 5, "g": 6, "h": 7,
        "а": 0, "б": 1, "в": 2, "г": 3, "д": 4, "е": 5,
        "1": 0, "2": 1, "3": 2, "4": 3, "5": 4,
    }
    qs = []
    for row in rows:
        cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
        if len(cells) < 3:
            continue
        first_lower = cells[0].lower()
        if any(h in first_lower for h in ("savol", "question", "вопрос", "№", "tartib", "nomer")):
            if len(cells[0]) < 30 and not any(ch in cells[0] for ch in ("?", "!", ":")):
                continue

        # Check if first column is just a row number
        if cells[0].isdigit() and len(cells[0]) <= 4 and len(cells) >= 4:
            q_raw = cells[1]
            options = list(cells[2:-1])
            raw_ans = cells[-1]
        else:
            q_raw = cells[0]
            options = list(cells[1:-1])
            raw_ans = cells[-1]

        if len(options) < 2:
            continue

        q_text = re.sub(r"^\d+\s*[.)]\s*", "", q_raw).strip()
        q_text = re.sub(r"^(savol|question|вопрос)\s*[:.\-)]*\s*", "", q_text, flags=re.I).strip()

        ans_idx = None
        ans_clean = re.sub(
            r"^(to[ʻ'’`]?g[ʻ'’`]?ri\s*javob|javob[a-z]*|answer|ответ|прав[a-z]*\s*ответ)\s*[:.\-)]*\s*",
            "",
            raw_ans,
            flags=re.I,
        ).strip().lower()

        if ans_clean in letters_map and letters_map[ans_clean] < len(options):
            ans_idx = letters_map[ans_clean]
        else:
            for oi, opt in enumerate(options):
                opt_clean = re.sub(r"^[a-hа-д0-9]\s*[).:]\s*", "", opt, flags=re.I).strip().lower()
                if raw_ans.lower() == opt.lower() or raw_ans.lower() == opt_clean:
                    ans_idx = oi
                    break
        if ans_idx is None:
            for oi, opt in enumerate(options):
                if opt.startswith("+") or opt.startswith("*"):
                    ans_idx = oi
                    break

        cleaned_opts = []
        for opt in options:
            clean_opt = re.sub(r"^[\+\*\-]?\s*[a-hа-дA-HА-Д0-9]?\s*[).:]\s*", "", opt).strip()
            if not clean_opt:
                clean_opt = opt
            cleaned_opts.append(clean_opt)

        if ans_idx is not None and 2 <= len(cleaned_opts) <= 8 and len(q_text) >= 2:
            qs.append({
                "q": q_text,
                "options": cleaned_opts,
                "answer": ans_idx,
            })
    return qs[:500]


def extract_excel(path) -> list[dict]:
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = []
    lines = []
    for r in ws.iter_rows(values_only=True):
        cells = [str(c).strip() for c in r if c is not None and str(c).strip()]
        if cells:
            rows.append(cells)
            for c in cells:
                lines.append(c)
    qs = parse_table_questions(rows)
    if not qs:
        qs = parse_questions(lines)
    return qs


def extract_csv(path) -> list[dict]:
    p = Path(path)
    raw = p.read_bytes()
    for enc in ("utf-8-sig", "utf-8", "cp1251", "latin-1"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        text = raw.decode("utf-8", errors="ignore")

    best_qs = []
    for delim in (",", ";", "\t", "|"):
        try:
            reader = csv.reader(text.splitlines(), delimiter=delim)
            rows = [[c.strip() for c in r if c.strip()] for r in reader]
            rows = [r for r in rows if r]
            qs = parse_table_questions(rows)
            if len(qs) > len(best_qs):
                best_qs = qs
        except Exception:
            continue
    if best_qs:
        return best_qs
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    return parse_questions(lines)


def extract_docx_questions(path) -> list[dict]:
    lines = extract_docx(path)
    qs = parse_questions(lines)
    if qs:
        return qs
    # If standard line parser didn't find questions, check if docx has tables with table parser
    try:
        doc = Document(path)
        rows = []
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    rows.append(cells)
        if rows:
            t_qs = parse_table_questions(rows)
            if t_qs:
                return t_qs
    except Exception:
        pass
    return []


def extract_questions_from_file(path: str | Path) -> list[dict]:
    p = Path(path)
    ext = p.suffix.lower()
    if ext == ".docx":
        return extract_docx_questions(p)
    elif ext == ".txt":
        lines = extract_txt(p)
        return parse_questions(lines)
    elif ext in (".xlsx", ".xls"):
        return extract_excel(p)
    elif ext == ".csv":
        return extract_csv(p)
    else:
        try:
            lines = extract_txt(p)
            return parse_questions(lines)
        except Exception:
            return []


def extract_questions_from_text(text: str) -> list[dict]:
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    return parse_questions(lines)


def letters_for(n):
    return ["A", "B", "C", "D", "E", "F"][:n]
