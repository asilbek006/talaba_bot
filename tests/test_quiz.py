from services import quiz


def test_uz_numeric():
    lines = [
        "1. Toshkent qaysi davlat poytaxti?",
        "A) Buxoro",
        "B) Toshkent",
        "C) Samarqand",
        "D) Qarshi",
        "To'g'ri javob: B",
    ]
    qs = quiz.parse_questions(lines)
    assert len(qs) == 1
    assert qs[0]["answer"] == 1
    assert len(qs[0]["options"]) == 4
    assert qs[0]["q"].startswith("Toshkent")


def test_savol_format():
    lines = [
        "Savol: 2+2 nechi?",
        "A) 3",
        "B) 4",
        "C) 5",
        "D) 6",
        "To'g'ri javob: B",
    ]
    qs = quiz.parse_questions(lines)
    assert len(qs) == 1
    assert qs[0]["q"].startswith("2+2")
    assert qs[0]["answer"] == 1


def test_ru_vopros_otvet():
    lines = [
        "Вопрос: Столица Узбекистана?",
        "A) Самарканд",
        "B) Ташкент",
        "C) Бухара",
        "D) Андижан",
        "Ответ: B",
    ]
    qs = quiz.parse_questions(lines)
    assert len(qs) == 1
    assert qs[0]["answer"] == 1
    assert qs[0]["q"].startswith("Столица")


def test_en_answer():
    lines = [
        "Question 1: Capital of France?",
        "A) Berlin",
        "B) Paris",
        "C) Madrid",
        "D) Rome",
        "Answer: B",
    ]
    qs = quiz.parse_questions(lines)
    assert len(qs) == 1
    assert qs[0]["answer"] == 1


def test_multiple_questions():
    lines = [
        "1. Savol 1",
        "A) a1",
        "B) b1",
        "C) c1",
        "D) d1",
        "Javob: A",
        "2. Savol 2",
        "A) a2",
        "B) b2",
        "C) c2",
        "D) d2",
        "To'g'ri javob: C",
    ]
    qs = quiz.parse_questions(lines)
    assert len(qs) == 2
    assert [q["answer"] for q in qs] == [0, 2]


def test_no_answer_missing():
    lines = ["1. Savol", "A) a", "B) b", "C) c", "D) d"]
    qs = quiz.parse_questions(lines)
    assert qs == []


def test_few_options_filtered():
    lines = ["1. Savol", "A) a", "Javob: A"]
    qs = quiz.parse_questions(lines)
    assert len(qs) == 0


def test_extract_docx(tmp_path):
    from docx import Document

    d = Document()
    d.add_paragraph("Savol:")
    d.add_paragraph("A) x")
    p = tmp_path / "q.docx"
    d.save(p)
    lines = quiz.extract_docx(p)
    assert lines == ["Savol:", "A) x"]
