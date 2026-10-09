from docx import Document
from openpyxl import load_workbook
from pptx import Presentation

from services import documents as docs


def test_build_referat(tmp_path):
    data = {
        "title": "Elektron tijorat",
        "sections": [
            {"heading": "Kirish", "paragraphs": ["Birinchi paragraf.", "Ikkinchi paragraf."]},
            {"heading": "Asosiy qism", "paragraphs": ["Uchinchi paragraf."]},
        ],
    }
    out = docs.build_referat(data, "uz", tmp_path / "ref.docx")
    assert out.exists()
    doc = Document(out)
    texts = [p.text for p in doc.paragraphs]
    assert any("Elektron tijorat" in x for x in texts)
    assert any("Birinchi paragraf" in x for x in texts)
    assert any("Kirish" in x for x in texts)


def test_build_referat_ru(tmp_path):
    data = {"title": "Тема", "sections": [{"heading": "Введение", "paragraphs": ["Текст."]}]}
    out = docs.build_referat(data, "ru", tmp_path / "ref_ru.docx")
    assert out.exists()
    texts = [p.text for p in Document(out).paragraphs]
    assert any("РЕФЕРАТ" in x for x in texts)


def test_build_pptx_slide_count(tmp_path):
    data = {
        "title": "AI",
        "subtitle": "Ta'limda",
        "slides": [
            {"title": "Kirish", "bullets": ["b1", "b2", "b3"], "notes": "n1"},
            {"title": "Xulosa", "bullets": ["b4", "b5"], "notes": "n2"},
        ],
    }
    out = docs.build_pptx(data, "uz", tmp_path / "p.pptx")
    assert out.exists()
    prs = Presentation(out)
    assert len(prs.slides) == 2


def test_build_pptx_notes(tmp_path):
    data = {"title": "T", "slides": [{"title": "S", "bullets": ["b"], "notes": "hop notes"}]}
    out = docs.build_pptx(data, "uz", tmp_path / "p2.pptx")
    assert out.exists()


def test_build_pptx_all_bullets_render(tmp_path):
    data = {
        "title": "T",
        "subtitle": "s",
        "slides": [
            {"title": "Muqova", "bullets": [], "notes": ""},
            {"title": "S2", "bullets": ["birinchi", "ikkinchi", "uchinchi"], "notes": "n"},
        ],
    }
    out = docs.build_pptx(data, "uz", tmp_path / "p3.pptx")
    prs = Presentation(out)
    slide = prs.slides[1]
    joined = "\n".join(s.text_frame.text for s in slide.shapes if s.has_text_frame)
    for w in ("birinchi", "ikkinchi", "uchinchi"):
        assert w in joined, f"bullet '{w}' render bo'lmadi"


def test_build_xlsx_rows_and_headers(tmp_path):
    data = {
        "title": "Aholi",
        "headers": ["Viloyat", "Aholi"],
        "rows": [["Toshkent", 2900000], ["Samarqand", 3900000]],
        "note": "taxminiy",
    }
    out = docs.build_xlsx(data, "uz", tmp_path / "t.xlsx")
    assert out.exists()
    wb = load_workbook(out)
    ws = wb.active
    assert ws.cell(row=2, column=1).value == "Viloyat"
    assert ws.cell(row=3, column=1).value == "Toshkent"


def test_build_xlsx_empty_rows(tmp_path):
    data = {"title": "Bo'sh", "headers": ["№", "Ma'lumot"], "rows": [], "note": ""}
    out = docs.build_xlsx(data, "uz", tmp_path / "t2.xlsx")
    assert out.exists()


def test_build_test_docx(tmp_path):
    qs = [
        {"q": "Savol 1?", "options": ["A v", "B v", "C v", "D v"], "answer": 1},
        {"q": "Savol 2?", "options": ["A v", "B v", "C v", "D v"], "answer": 3},
    ]
    out = docs.build_test_docx(qs, "uz", tmp_path / "test.docx")
    assert out.exists()
    texts = [p.text for p in Document(out).paragraphs]
    assert any("TEST SAVOLLARI" in x for x in texts)
    assert any("Toʻgʻri javob: B" in x for x in texts)
    assert any("Toʻgʻri javob: D" in x for x in texts)


def test_build_test_docx_bad_answer(tmp_path):
    qs = [{"q": "Q?", "options": ["A", "B", "C", "D"], "answer": 99}]
    out = docs.build_test_docx(qs, "uz", tmp_path / "t.docx")
    texts = [p.text for p in Document(out).paragraphs]
    assert any("Toʻgʻri javob: ?" in x for x in texts)


def test_build_text_docx(tmp_path):
    paras = ["Oddiy paragraf.", "## Sarlavha", "Yana matn"]
    out = docs.build_text_docx(paras, "Sarlavha", tmp_path / "txt.docx")
    assert out.exists()


def test_build_test_docx_state_format(tmp_path):
    questions = [{"q": "Savol?", "options": ["A1", "B1", "C1", "D1"], "answer": 0}]
    out = docs.build_test_docx(questions, "uz", tmp_path / "t.docx")
    doc = Document(out)
    sec = doc.sections[0]
    from docx.shared import Cm, Pt

    assert abs(sec.left_margin - Cm(3)) < 5000
    assert abs(sec.right_margin - Cm(1.5)) < 5000
    normal = doc.styles["Normal"]
    assert normal.font.size == Pt(14)
    q_para = next(p for p in doc.paragraphs if p.text.startswith("1."))
    assert q_para.paragraph_format.line_spacing == 1.5


def test_build_text_docx_state_format(tmp_path):
    out = docs.build_text_docx(["Matn."], "Sarlavha", tmp_path / "t2.docx")
    doc = Document(out)
    from docx.shared import Cm, Pt

    assert doc.styles["Normal"].font.size == Pt(14)
    body = next(p for p in doc.paragraphs if p.text == "Matn.")
    assert body.paragraph_format.line_spacing == 1.5
    assert abs(body.paragraph_format.first_line_indent - Cm(1.25)) < 5000


def test_build_pptx_state_format_and_image(tmp_path):
    from io import BytesIO

    from PIL import Image

    buf = BytesIO()
    Image.new("RGB", (320, 180), (200, 60, 60)).save(buf, format="PNG")
    data = {
        "title": "Mavzu",
        "slides": [
            {"title": "Muqova", "bullets": [], "notes": ""},
            {
                "title": "Sahna",
                "bullets": ["Birinchi", "Ikkinchi"],
                "notes": "N",
                "image_bytes": buf.getvalue(),
            },
        ],
    }
    out = docs.build_pptx(data, "uz", tmp_path / "s.pptx")
    prs = Presentation(out)
    assert len(prs.slides) == 2
    body_slide = prs.slides[1]
    sizes, spacings, pictures = set(), set(), 0
    for sh in body_slide.shapes:
        if sh.shape_type == 13:
            pictures += 1
        if sh.has_text_frame:
            for para in sh.text_frame.paragraphs:
                for r in para.runs:
                    if r.font.size:
                        sizes.add(r.font.size.pt)
                if para.line_spacing:
                    spacings.add(para.line_spacing)
    assert pictures == 1
    assert 14.0 in sizes
    assert 1.5 in spacings


def test_xlsx_ragged_rows_are_padded(tmp_path):
    out = docs.build_xlsx(
        {"title": "Ragged", "headers": ["A", "B"], "rows": [["one"]]},
        "en",
        tmp_path / "ragged.xlsx",
    )
    sheet = load_workbook(out).active
    assert sheet["A3"].value == "one"
    assert sheet["B3"].value is None


def test_xlsx_ai_strings_are_not_executable_formulas(tmp_path):
    data = {
        "title": "=1+1",
        "headers": ["=SUM(A1)", "Value"],
        "rows": [['=HYPERLINK("https://example.com")', 10]],
        "note": "=1+1",
    }
    sheet = load_workbook(docs.build_xlsx(data, "en", tmp_path / "safe.xlsx")).active
    for cell in ("A1", "A2", "A3", "A5"):
        assert sheet[cell].data_type == "s"


def test_single_page_report_does_not_force_a_cover_break(tmp_path):
    data = {
        "title": "Topic",
        "include_cover": False,
        "sections": [{"heading": "Intro", "paragraphs": ["text"]}],
    }
    document = Document(docs.build_referat(data, "en", tmp_path / "short.docx"))
    assert 'w:type="page"' not in document._element.xml


def test_generated_test_labels_follow_language(tmp_path):
    questions = [{"q": "Question", "options": ["a", "b", "c", "d"], "answer": 1}]
    for lang, label in [("en", "Correct answer"), ("ru", "Правильный ответ")]:
        document = Document(docs.build_test_docx(questions, lang, tmp_path / f"{lang}.docx"))
        assert any(label in p.text for p in document.paragraphs)
