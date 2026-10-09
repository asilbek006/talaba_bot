import subprocess
from types import SimpleNamespace as NS
from zipfile import ZipFile

import pytest
from PIL import Image

import config
from services import convert
from services.validation import validate_docx
from tests.test_pdf import make_pdf


def test_pdf_fallback_has_deadline_and_removes_failed_workspace(tmp_path, monkeypatch):
    workspace = tmp_path / "convert"
    workspace.mkdir()
    monkeypatch.setattr(convert, "TMP", workspace)
    source = make_pdf(tmp_path / "source.pdf")
    calls = []

    def run(argv, **kwargs):
        calls.append((argv, kwargs))
        if argv[0] == "soffice":
            return NS(returncode=1)
        raise subprocess.TimeoutExpired(argv, kwargs["timeout"])

    monkeypatch.setattr(convert.subprocess, "run", run)
    with pytest.raises(subprocess.TimeoutExpired):
        convert.pdf_to_docx_sync(source)
    assert len(calls) == 2
    assert calls[1][0][1:3] == ["-m", "services.convert_worker"]
    assert calls[1][1]["timeout"] == 180
    assert list(workspace.iterdir()) == []


def test_fallback_worker_converts_real_pdf(tmp_path):
    import shutil

    from docx import Document
    from pypdf import PdfReader

    doc = Document()
    doc.add_paragraph("Isolated conversion worker sample")
    source = tmp_path / "sample.docx"
    doc.save(source)
    pdf = convert.docx_to_pdf_sync(source)
    saved = tmp_path / "source.pdf"
    shutil.copyfile(pdf, saved)
    convert.cleanup_result(pdf)
    assert len(PdfReader(saved).pages) == 1
    output = tmp_path / "converted.docx"
    result = subprocess.run(
        [convert.sys.executable, "-m", "services.convert_worker", str(saved), str(output)],
        cwd=config.BASE_DIR,
        capture_output=True,
        text=True,
        timeout=180,
    )
    assert result.returncode == 0, result.stderr
    assert "sample" in "\n".join(p.text for p in Document(output).paragraphs)


def test_invalid_docx_rejected_before_libreoffice(tmp_path, monkeypatch):
    path = tmp_path / "fake.docx"
    path.write_bytes(b"not a zip")
    called = []
    monkeypatch.setattr(convert, "_run_soffice", lambda *args: called.append(args))
    with pytest.raises(ValueError):
        convert.docx_to_pdf_sync(path)
    assert not called
    with ZipFile(path, "w") as archive:
        archive.writestr("irrelevant.txt", "not a Word document")
    with pytest.raises(ValueError):
        validate_docx(path)


def test_oversized_pdf_and_image_are_rejected(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "MAX_PDF_PAGES", 1)
    with pytest.raises(ValueError):
        convert.pdf_to_docx_sync(make_pdf(tmp_path / "long.pdf", 2))
    path = tmp_path / "wide.png"
    with Image.new("RGB", (8001, 2)) as image:
        image.save(path)
    output = tmp_path / "result.pdf"
    with pytest.raises(ValueError):
        convert.images_to_pdf_sync([path], output)
    assert not output.exists()
